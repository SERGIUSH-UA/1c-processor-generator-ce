"""
Tests for Excel -> MXL conversion formatting (v2.78.0+)

Before v2.78.0 the converter silently dropped background colors, font colors,
border styles, number formats and images: the CLI reported success and the
resulting template simply looked nothing like the source.

Tag names and enum values used here were taken from real 1C templates, not invented.
"""

import base64
import io
import re
from pathlib import Path
import sys

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

openpyxl = pytest.importorskip("openpyxl", reason="Excel conversion requires openpyxl")

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import importlib
excel_to_mxl = importlib.import_module("1c_processor_generator.pro.excel_to_mxl")
mxl_builder = importlib.import_module("1c_processor_generator.pro.mxl_builder")

ExcelToMXLConverter = excel_to_mxl.ExcelToMXLConverter
convert_number_format = excel_to_mxl.convert_number_format
MXLFormat = mxl_builder.MXLFormat

# 1x1 transparent PNG
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _convert(wb, tmp_path) -> tuple:
    """Save a workbook and convert it; return (xml, warnings)."""
    path = tmp_path / "template.xlsx"
    wb.save(path)

    converter = ExcelToMXLConverter(languages=["ru", "uk"])
    xml = converter.convert(str(path))
    return xml, converter.warnings


class TestColors:
    def test_background_and_font_colors_are_emitted(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Заголовок"
        ws["A1"].fill = PatternFill("solid", fgColor="FFDDFFDA")
        ws["A1"].font = Font(color="FF1F497D")

        xml, _ = _convert(wb, tmp_path)

        assert "<backColor>#DDFFDA</backColor>" in xml
        assert "<textColor>#1F497D</textColor>" in xml

    def test_black_text_is_not_emitted(self, tmp_path):
        """Black is the 1C default - emitting it would bloat the format registry"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Звичайний текст"
        ws["A1"].font = Font(color="FF000000")

        xml, _ = _convert(wb, tmp_path)

        assert "<textColor>#000000</textColor>" not in xml

    def test_cell_without_fill_has_no_background(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Без заливки"

        xml, _ = _convert(wb, tmp_path)

        assert "<backColor>" not in xml


class TestBorders:
    def test_border_styles_map_to_distinct_lines(self, tmp_path):
        """thick/double must not both collapse into the default solid line"""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Рамка"
        ws["A1"].border = Border(
            left=Side(style="thick", color="FFC0504D"),
            top=Side(style="double"),
        )

        xml, _ = _convert(wb, tmp_path)

        assert '<line width="3"' in xml     # thick
        assert "Double" in xml
        assert "<borderColor>#C0504D</borderColor>" in xml

    def test_cell_without_border_registers_no_color(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Без рамки"

        xml, _ = _convert(wb, tmp_path)

        assert "<borderColor>" not in xml


class TestNumberFormats:
    @pytest.mark.parametrize("excel_format,expected", [
        ("0.00", "ЧЦ=15; ЧДЦ=2"),
        ("0", "ЧЦ=15; ЧДЦ=0"),
        ("#,##0.00", "ЧЦ=15; ЧДЦ=2; ЧРГ=' '"),
        ("dd.mm.yyyy", "ДФ=dd.MM.yyyy"),
    ])
    def test_supported_formats(self, excel_format, expected):
        assert convert_number_format(excel_format) == expected

    @pytest.mark.parametrize("excel_format", ["General", "@", "0%", None, ""])
    def test_unsupported_formats_return_none(self, excel_format):
        assert convert_number_format(excel_format) is None

    def test_format_reaches_xml_as_localized_element(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "{Сумма}"
        ws["A1"].number_format = "#,##0.00"

        xml, _ = _convert(wb, tmp_path)

        assert "ЧЦ=15; ЧДЦ=2; ЧРГ=' '" in xml
        assert "<v8:lang>ru</v8:lang>" in xml
        assert "<v8:lang>uk</v8:lang>" in xml

    def test_untranslatable_format_warns(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "{Процент}"
        ws["A1"].number_format = "0%"

        _, warnings = _convert(wb, tmp_path)

        assert any("0%" in w for w in warnings)


class TestImages:
    def test_image_becomes_drawing_and_picture(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "З логотипом"

        image = XLImage(io.BytesIO(TINY_PNG))
        image.width, image.height = 120, 40
        ws.add_image(image, "C1")

        xml, warnings = _convert(wb, tmp_path)

        assert "<drawingType>Picture</drawingType>" in xml
        # 1-based reference: <pictureIndex>0</pictureIndex> means "no picture" to 1C,
        # which then drops the registry entry and loses the image
        assert "<pictureIndex>1</pictureIndex>" in xml
        assert "<index>0</index>" in xml
        assert base64.b64encode(TINY_PNG).decode("ascii") in xml
        # The logo is carried over, so it must NOT be reported as lost
        assert not any("зображ" in w.lower() for w in warnings)

    def test_no_images_no_drawings(self, tmp_path):
        wb = Workbook()
        wb.active["A1"] = "Без картинок"

        xml, _ = _convert(wb, tmp_path)

        assert "<drawing>" not in xml
        assert "<picture>" not in xml


class TestFormatRegistry:
    def test_colors_are_part_of_the_dedup_key(self):
        """
        Without colors in to_key(), two differently colored cells would share
        one registry entry and render identically.
        """
        red = MXLFormat(back_color="#FF0000")
        blue = MXLFormat(back_color="#0000FF")

        assert red.to_key() != blue.to_key()

    def test_identical_formats_share_a_key(self):
        assert MXLFormat(back_color="#FF0000").to_key() == MXLFormat(back_color="#FF0000").to_key()


class TestOutputValidity:
    def test_output_is_well_formed_xml(self, tmp_path):
        import xml.etree.ElementTree as ET

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Тест {Параметр}"
        ws["A1"].fill = PatternFill("solid", fgColor="FFDDFFDA")
        ws["A1"].alignment = Alignment(horizontal="center", indent=2)
        ws["A1"].border = Border(bottom=Side(style="medium"))

        xml, _ = _convert(wb, tmp_path)

        ET.fromstring(xml.encode("utf-8"))  # raises on malformed XML

    def test_indent_is_emitted(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Відступ"
        ws["A1"].alignment = Alignment(indent=2)

        xml, _ = _convert(wb, tmp_path)

        assert "<indent>2</indent>" in xml
